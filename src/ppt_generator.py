from datetime import datetime, timezone
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

_MAX_CARRYOVER_SHOWN = 15


def build(sprint_name: str, stats: dict, narrative: dict) -> bytes:
    prs = Presentation()

    _add_title_slide(prs, sprint_name)
    _add_overview_slide(prs, stats)
    _add_status_breakdown_slide(prs, stats)
    _add_highlights_risks_slide(prs, narrative)
    _add_next_steps_slide(prs, stats, narrative)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_title_slide(prs: Presentation, sprint_name: str) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = sprint_name
    date_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    slide.placeholders[1].text = f"Sprint Summary — generated {date_str}"


def _set_bullets(text_frame, lines: list[str]) -> None:
    text_frame.text = lines[0]
    for line in lines[1:]:
        p = text_frame.add_paragraph()
        p.text = line


def _add_overview_slide(prs: Presentation, stats: dict) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Overview"

    committed = stats["committed_points"]
    completed = stats["completed_points"]
    pct = (completed / committed * 100) if committed else 0.0

    lines = [
        f"Committed points: {committed:g}",
        f"Completed points: {completed:g}",
        f"% complete: {pct:.0f}%",
        f"Total issues: {stats['total_issues']}",
    ]
    for status, count in stats["status_counts"].items():
        lines.append(f"{status}: {count}")

    _set_bullets(slide.placeholders[1].text_frame, lines)


def _add_status_breakdown_slide(prs: Presentation, stats: dict) -> None:
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Status Breakdown"

    chart_data = CategoryChartData()
    chart_data.categories = list(stats["status_counts"].keys())
    chart_data.add_series("Issues", list(stats["status_counts"].values()))

    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1), Inches(1.5), Inches(8), Inches(5),
        chart_data,
    )


def _add_highlights_risks_slide(prs: Presentation, narrative: dict) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Highlights & Risks"

    highlights = narrative.get("highlights") or ["None reported."]
    risks = narrative.get("risks") or ["None reported."]

    lines = ["Highlights:"]
    lines += [f"• {h}" for h in highlights]
    lines.append("Risks:")
    lines += [f"• {r}" for r in risks]

    _set_bullets(slide.placeholders[1].text_frame, lines)


def _add_next_steps_slide(prs: Presentation, stats: dict, narrative: dict) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Next Steps"

    next_steps = narrative.get("next_steps") or ["None reported."]
    carryover = stats["carryover_keys"]
    shown = carryover[:_MAX_CARRYOVER_SHOWN]

    lines = list(next_steps)
    lines.append("Carryover:")
    lines += shown
    if len(carryover) > _MAX_CARRYOVER_SHOWN:
        lines.append(f"+{len(carryover) - _MAX_CARRYOVER_SHOWN} more")

    _set_bullets(slide.placeholders[1].text_frame, lines)
